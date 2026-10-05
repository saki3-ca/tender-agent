"""
Text extraction from tender documents (PDF, DOCX, XLSX and scanned notices).

Only the first few pages are read: the notice details (subject, publication date,
submission deadline, scope) are on the first pages, and reading whole bidding documents
would slow the run without improving the result. OCR (English + Bangla) is used for scanned
pages and image notices when Tesseract is installed (it is installed in GitHub Actions).
"""

import io
from functools import lru_cache
from typing import Optional
from urllib.parse import urlparse

from app.utils.logging import logger

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None
try:
    import docx
except ImportError:  # pragma: no cover
    docx = None
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None
try:
    import pytesseract
    from PIL import Image
except ImportError:  # pragma: no cover
    pytesseract = None
    Image = None


@lru_cache(maxsize=1)
def ocr_available() -> bool:
    if not pytesseract:
        return False
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr(image) -> str:
    try:
        return pytesseract.image_to_string(image, lang="eng+ben")
    except Exception:
        try:
            return pytesseract.image_to_string(image, lang="eng")
        except Exception as e:
            logger.warning(f"OCR failed: {e}")
            return ""


def _doc_kind(url: str, content_type: str, content: bytes) -> str:
    path = urlparse(url).path.lower()
    if content[:5] == b"%PDF-" or path.endswith(".pdf") or "pdf" in content_type:
        return "pdf"
    if path.endswith(".docx") or "wordprocessingml" in content_type:
        return "docx"
    if path.endswith((".xlsx", ".xls")) or "spreadsheet" in content_type:
        return "xlsx"
    if path.endswith((".jpg", ".jpeg", ".png")) or content_type.startswith("image/"):
        return "image"
    return "other"


def extract_document_text(content: bytes, url: str = "", content_type: str = "",
                          max_pages: int = 4, max_chars: int = 20000) -> str:
    """Best-effort text of a tender document; empty string if nothing can be read."""
    kind = _doc_kind(url, content_type, content)
    try:
        if kind == "pdf" and pdfplumber:
            parts = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for idx, page in enumerate(pdf.pages[:max_pages]):
                    text = page.extract_text() or ""
                    if len(text.strip()) < 40 and idx < 2 and ocr_available():
                        text = _ocr(page.to_image(resolution=200).original) or text
                    parts.append(text)
            return "\n".join(parts)[:max_chars]
        if kind == "docx" and docx:
            d = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in d.paragraphs if p.text.strip())[:max_chars]
        if kind == "xlsx" and openpyxl:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            lines = []
            for ws in wb.worksheets[:2]:
                for row in ws.iter_rows(values_only=True, max_row=200):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines)[:max_chars]
        if kind == "image" and Image and ocr_available():
            return _ocr(Image.open(io.BytesIO(content)))[:max_chars]
    except Exception as e:  # noqa: BLE001 - unreadable document is not fatal
        logger.warning(f"Could not read document {url}: {type(e).__name__}: {e}")
    return ""


def is_readable_document(url: Optional[str]) -> bool:
    if not url:
        return False
    path = urlparse(url).path.lower()
    if path.endswith((".jpg", ".jpeg", ".png")):
        return ocr_available()
    return not path.endswith(".zip")
